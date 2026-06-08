import shutil
import openpyxl
from openpyxl.styles import PatternFill, Font

_HIGHLIGHT = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
_BOLD = Font(bold=True)


def _find_header_col(ws, name: str) -> int | None:
    for cell in ws[1]:
        if cell.value and str(cell.value).strip() == name:
            return cell.column
    return None


def _add_refined_columns(ws, existing_col_names: list[str]) -> dict[str, int]:
    header_row = ws[1]
    last_col = max((c.column for c in header_row if c.value), default=0)

    col_map = {}
    for name in existing_col_names:
        refined_name = name + "_refined"
        # Check if already exists
        existing = _find_header_col(ws, refined_name)
        if existing:
            col_map[name] = existing
        else:
            last_col += 1
            col_map[name] = last_col
            cell = ws.cell(row=1, column=last_col, value=refined_name)
            cell.font = _BOLD
            cell.fill = _HIGHLIGHT

    return col_map


def write_output(
    input_path: str,
    manga_results: dict[int, dict],  # row_idx -> {th, en, jp}
    novel_results: dict[int, dict],
    output_path: str,
) -> None:
    shutil.copy2(input_path, output_path)
    wb = openpyxl.load_workbook(output_path)

    # --- Manga sheet ---
    ws_manga = wb["manga_translation_master_FINAL"]
    col_map_manga = _add_refined_columns(ws_manga, ["Text_TH", "Text_EN", "Text_JP"])

    for row_idx, refined in manga_results.items():
        ws_manga.cell(row=row_idx, column=col_map_manga["Text_TH"], value=refined.get("th", ""))
        ws_manga.cell(row=row_idx, column=col_map_manga["Text_EN"], value=refined.get("en", ""))
        ws_manga.cell(row=row_idx, column=col_map_manga["Text_JP"], value=refined.get("jp", ""))

    # --- Novel sheet ---
    ws_novel = wb["novel_translation_master_FINAL"]
    col_map_novel = _add_refined_columns(ws_novel, ["Text_TH", "Text_EN", "Text_JP"])

    for row_idx, refined in novel_results.items():
        ws_novel.cell(row=row_idx, column=col_map_novel["Text_TH"], value=refined.get("th", ""))
        ws_novel.cell(row=row_idx, column=col_map_novel["Text_EN"], value=refined.get("en", ""))
        ws_novel.cell(row=row_idx, column=col_map_novel["Text_JP"], value=refined.get("jp", ""))

    wb.save(output_path)
    wb.close()
